"""Standardized workbook export with LIVE formulas for the valuation tail.

One sheet per scenario: yellow input cells (WACC, g, stub, shares, bridge
include flags, exit multiple), projection rows as values (they are the output
of the agent's archetype engine), and formulas from NOPAT/FCFF through
discounting, terminal value, EV, the bridge and value per share. The formulas
mirror dcfkit.engine exactly; `verify_export()` recalculates through
LibreOffice and compares.
"""
from __future__ import annotations

import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

YELLOW = PatternFill("solid", fgColor="FFF2CC")
HEAD = PatternFill("solid", fgColor="DDEBF7")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
MUTED = Font(italic=True, color="808080")
THIN = Border(top=Side(style="thin"))
NUM = '#,##0.0;[Red]-#,##0.0'
NUM2 = '#,##0.00;[Red]-#,##0.00'
PCT = '0.0%'
FIRST_COL = 3  # periods start in column C


def _sheet_name(s, used):
    base = re.sub(r"[\[\]\*\?/\\:]", "", s)[:28] or "Scenario"
    name, n = base, 2
    while name in used:
        name = f"{base[:25]}_{n}"
        n += 1
    used.add(name)
    return name


def _q(sheet, cell):
    return f"'{sheet}'!{cell}"


class _W:
    def __init__(self, ws):
        self.ws = ws
        self.r = 1

    def put(self, row, col, value, *, fmt=None, font=None, fill=None, border=None):
        c = self.ws.cell(row=row, column=col, value=value)
        if fmt:
            c.number_format = fmt
        if font:
            c.font = font
        if fill:
            c.fill = fill
        if border:
            c.border = border
        return c

    def label(self, text, value=None, *, fmt=None, inp=False, note=None, font=None):
        self.put(self.r, 1, text, font=font)
        c = None
        if value is not None:
            c = self.put(self.r, 2, value, fmt=fmt, fill=YELLOW if inp else None, font=font)
        if note:
            self.put(self.r, 3, note, font=MUTED)
        addr = f"B{self.r}"
        self.r += 1
        return addr


def _explicit(rows):
    return [r for r in rows or [] if not r.get("is_actual")]


def _write_scenario(wb, name, scn, dcf):
    ws = wb.create_sheet(name)
    w = _W(ws)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 14
    w.put(1, 1, f"{dcf['provenance']['ticker']} — {scn.get('label', scn['id'])} scenario", font=TITLE)
    w.put(2, 1, "Yellow cells are inputs. Projection rows are values from the model's own engine; "
                "everything from NOPAT down is a live formula.", font=MUTED)
    w.r = 4
    w.put(w.r, 1, "Inputs", font=BOLD, fill=HEAD)
    w.r += 1
    disc = scn["discounting"]
    wacc = w.label("WACC", disc["wacc"], fmt=PCT, inp=True)
    conv = w.label("Discount convention (end/mid)", disc.get("convention", "end"), inp=True)
    stub = w.label("Stub fraction (share of period 1 ahead)", disc.get("stub_fraction", 1.0), fmt="0.000", inp=True)
    stub_f = w.label("Stub FCFF override (blank = FCFF × stub)", disc.get("stub_fcff"), fmt=NUM, inp=True)
    if disc.get("stub_fcff") is None:
        ws[stub_f] = None
        ws[stub_f].fill = YELLOW
    tvt = w.label("TV timing (end_last/mid_last)", disc.get("tv_timing", "end_last"), inp=True)
    sh = scn["shares"]
    shares = w.label("Shares used (M, ordinary)", sh["used"], fmt=NUM2, inp=True, note=sh.get("method"))
    proceeds = w.label("Issuance proceeds ($M)", sh.get("proceeds") or 0.0, fmt=NUM, inp=True)
    ratio = w.label("Ordinary shares per listed unit", (dcf.get("listing") or {}).get("ordinary_per_listed_unit", 1.0),
                    fmt="0.000", inp=True)
    fx = w.label("FX reporting → quote currency", (dcf.get("market") or {}).get("fx_spot", 1.0), fmt="0.0000", inp=True)
    price = w.label("Price per listed unit", dcf["market"]["price"], fmt=NUM2, inp=True,
                    note=str(dcf["market"].get("price_ts", "")))
    w.r += 1
    w.put(w.r, 1, "EV → equity bridge", font=BOLD, fill=HEAD)
    w.put(w.r, 2, "Equity effect", font=BOLD, fill=HEAD)
    w.put(w.r, 3, "Include (1/0)", font=BOLD, fill=HEAD)
    w.r += 1
    overrides = scn.get("bridge_overrides") or {}
    b0 = w.r
    for item in dcf.get("bridge") or []:
        inc = overrides.get(item["id"], item.get("included", True))
        w.put(w.r, 1, item["label"] + (" (debatable)" if item.get("debatable") else ""))
        w.put(w.r, 2, item["equity_effect"], fmt=NUM)
        w.put(w.r, 3, 1 if inc else 0, fill=YELLOW)
        w.r += 1
    b1 = w.r - 1
    bridge_total = f"B{w.r}"
    w.put(w.r, 1, "Bridge total", font=BOLD)
    w.put(w.r, 2, f"=SUMPRODUCT(B{b0}:B{b1},C{b0}:C{b1})" if b1 >= b0 else 0, fmt=NUM, font=BOLD)
    w.r += 2

    ent_values, eq_values, pv_tv_cells = [], [], []
    for comp in scn.get("components") or []:
        val, pv_tv = _write_component(w, comp, wacc, conv, stub, stub_f, tvt)
        (eq_values if comp.get("level") == "equity" else ent_values).append(val)
        if pv_tv and comp.get("level", "enterprise") != "equity":
            pv_tv_cells.append(pv_tv)
        w.r += 1

    w.put(w.r, 1, "Valuation", font=BOLD, fill=HEAD)
    w.r += 1
    ev = w.label("Enterprise value ($M)", "=" + ("+".join(ent_values) if ent_values else "0"), fmt=NUM, font=BOLD)
    w.label("TV share of EV", f"=IF({ev}>0,({'+'.join(pv_tv_cells) or '0'})/{ev},\"\")", fmt=PCT)
    w.label("+ Bridge total", f"={bridge_total}", fmt=NUM)
    eqc = w.label("+ Equity-level components", "=" + ("+".join(eq_values) if eq_values else "0"), fmt=NUM)
    w.label("+ Issuance proceeds", f"={proceeds}", fmt=NUM)
    eq = w.label("Equity value ($M)", f"={ev}+{bridge_total}+{eqc}+{proceeds}", fmt=NUM, font=BOLD)
    vpo = w.label("Value per ordinary share (reporting ccy)", f"={eq}/{shares}", fmt=NUM2)
    raw = w.label("Value per listed unit (raw)", f"={vpo}*{ratio}*{fx}", fmt=NUM2)
    vps = w.label("VALUE PER LISTED UNIT", f"=MAX({raw},0)", fmt=NUM2, font=BOLD)
    w.label("Upside vs price", f"={vps}/{price}-1", fmt=PCT)
    ws.freeze_panes = "C4"
    return {"vps": vps, "ev": ev, "equity": eq}


def _write_component(w, comp, scen_wacc, conv, stub, stub_f, tvt):
    ws = w.ws
    kind = comp.get("kind", "dcf")
    w.put(w.r, 1, f"Component: {comp.get('label', comp['id'])} ({kind}, {comp.get('level', 'enterprise')})",
          font=BOLD, fill=HEAD)
    w.r += 1
    if kind == "fixed":
        return w.label("Value ($M)", comp.get("fixed_value") or 0.0, fmt=NUM, inp=True), None
    cd = comp.get("discounting") or {}
    wacc = w.label("WACC (component)", cd["wacc"] if cd.get("wacc") is not None else f"={scen_wacc}", fmt=PCT,
                   inp=cd.get("wacc") is not None, note="own rate" if cd.get("wacc") is not None else "= scenario WACC")
    prob = w.label("Probability of success", comp.get("probability") if comp.get("probability") is not None else 1.0,
                   fmt=PCT, inp=True)
    term = comp.get("terminal") or {"method": "none"}
    method = term.get("method", "none")
    w.label("Terminal method", method)
    g = w.label("Terminal growth g", term.get("g") or 0.0, fmt=PCT, inp=True)
    g0 = w.label("g at which the model's stated N+1 FCF applies", term.get("g") or 0.0, fmt=PCT)
    mult = w.label("Exit multiple", term.get("multiple"), fmt="0.0x", inp=method == "exit_multiple")
    rows = _explicit(comp.get("rows"))
    n = len(rows)
    hdr = w.r
    w.put(hdr, 1, "Period", font=BOLD)
    for j, r in enumerate(rows):
        col = FIRST_COL + j
        w.put(hdr, col, r["period"], font=BOLD)
        ws.column_dimensions[get_column_letter(col)].width = 12
    lines = {}
    order = [("revenue", "Revenue"), ("ebitda", "EBITDA"), ("ebit", "EBIT"), ("taxes", "Taxes"),
             ("nopat", "NOPAT"), ("da", "D&A"), ("capex", "Capex"), ("delta_nwc", "Change in NWC"),
             ("other_adj", "Other adjustments"), ("fcff", "FCFF")]
    full = all(all(r.get(k) is not None for k in ("ebit", "taxes", "da", "capex", "delta_nwc")) for r in rows)
    rr = hdr + 1
    for key, label in order:
        w.put(rr, 1, label, font=BOLD if key in ("fcff", "nopat") else None)
        lines[key] = rr
        rr += 1
    for j, r in enumerate(rows):
        col = get_column_letter(FIRST_COL + j)
        for key, _ in order:
            row = lines[key]
            if key == "nopat" and full:
                v = f"={col}{lines['ebit']}-{col}{lines['taxes']}"
            elif key == "fcff" and full:
                v = (f"={col}{lines['nopat']}+{col}{lines['da']}-{col}{lines['capex']}"
                     f"-{col}{lines['delta_nwc']}+{col}{lines['other_adj']}")
            else:
                v = r.get(key)
                if key == "other_adj" and v is None:
                    v = 0.0
            w.put(row, FIRST_COL + j, v, fmt=NUM)
    t_row, df_row, pv_row = rr, rr + 1, rr + 2
    w.put(t_row, 1, "Discount time (years)")
    w.put(df_row, 1, "Discount factor")
    w.put(pv_row, 1, "PV of FCFF", font=BOLD)
    for j in range(n):
        col = get_column_letter(FIRST_COL + j)
        if j == 0:
            t = f'=IF({conv}="mid",{stub}/2,{stub})'
            flow = f'IF(ISNUMBER({stub_f}),{stub_f},{col}{lines["fcff"]}*{stub})'
        else:
            t = f'=IF({conv}="mid",{stub}+{j}-0.5,{stub}+{j})'
            flow = f"{col}{lines['fcff']}"
        w.put(t_row, FIRST_COL + j, t, fmt="0.00")
        w.put(df_row, FIRST_COL + j, f"=1/(1+{wacc})^{col}{t_row}", fmt="0.0000")
        w.put(pv_row, FIRST_COL + j, f"={flow}*{col}{df_row}", fmt=NUM)
    w.r = pv_row + 2
    last = get_column_letter(FIRST_COL + n - 1) if n else None
    first = get_column_letter(FIRST_COL)
    sum_pv = w.label("Sum of PV (explicit)", f"=SUM({first}{pv_row}:{last}{pv_row})" if n else 0, fmt=NUM)
    if n and method in ("normalized_gordon", "gordon", "finite_annuity"):
        L = lambda k: f"{last}{lines[k]}"  # noqa: E731
        nwc = term.get("nwc_pct_of_revenue")
        nwc_term = f"-{nwc}*{L('revenue')}*{{g}}" if (nwc and rows[-1].get("revenue") is not None) else ""
        if method == "gordon":
            F = lambda gg: f"{L('fcff')}*(1+{gg})"  # noqa: E731
        else:
            F = lambda gg: f"({L('nopat')}-{gg}*{L('da')})*(1+{gg}){nwc_term.format(g=gg)}"  # noqa: E731
        stored = term.get("fcf_next")
        if stored is not None and full:
            num = f"={stored}+({F(g)})-({F(g0)})"
            note = "stated N+1 FCF, re-derived when g moves"
        elif stored is not None:
            num, note = stored, "stated N+1 FCF (static)"
        else:
            num, note = f"={F(g)}", "normalized N+1 FCF"
        fnext = w.label("Terminal numerator: FCF N+1", num, fmt=NUM, note=note)
        if method == "finite_annuity":
            yrs = term.get("finite_years") or 1
            tv_formula = f"={fnext}/({wacc}-{g})*(1-((1+{g})/(1+{wacc}))^{yrs})"
        else:
            tv_formula = f"={fnext}/({wacc}-{g})"
    elif n and method == "exit_multiple":
        metric = term.get("multiple_metric", "ebitda")
        base = f"{last}{lines.get(metric, lines['ebitda'])}"
        if term.get("metric_period") == "next":
            base = f"{base}*(1+{g})"
        tv_formula = f"={mult}*{base}"
    else:
        tv_formula = 0
    tv = w.label("Terminal value", tv_formula, fmt=NUM)
    if term.get("floor_at_zero"):
        ws[tv] = f"=MAX({str(tv_formula).lstrip('=')},0)"
    T = w.label("TV discount time", f'=IF({tvt}="mid_last",{last}{t_row},{stub}+{n}-1)' if n else 0, fmt="0.00")
    pv_tv = w.label("PV of terminal value", f"={tv}/(1+{wacc})^{T}" if n else 0, fmt=NUM)
    val = w.label("Component value ($M)", f"={prob}*({sum_pv}+{pv_tv})", fmt=NUM, font=BOLD)
    pv_tv_w = w.label("Component PV of TV (prob-weighted)", f"={prob}*{pv_tv}", fmt=NUM)
    return val, pv_tv_w


def export(dcf):
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    used = {"Summary"}
    out = {}
    for scn in dcf["scenarios"]:
        name = _sheet_name(scn.get("label") or scn["id"], used)
        out[scn["id"]] = (name, _write_scenario(wb, name, scn, dcf))
    w = _W(summary)
    prov = dcf["provenance"]
    w.put(1, 1, f"{prov['ticker']} — {prov.get('company', '')}", font=TITLE)
    w.put(2, 1, f"DCF Studio export · source {prov.get('producer')} {prov.get('cli') or ''} "
                f"{prov.get('model') or ''} · version {prov.get('version_id') or ''} · as of "
                f"{dcf['basis']['as_of']}", font=MUTED)
    w.r = 4
    for k, v in (("Price per listed unit", dcf["market"]["price"]), ("Price date", dcf["market"].get("price_ts")),
                 ("Quote currency", (dcf.get("listing") or {}).get("quote_currency")),
                 ("Reporting currency", dcf["basis"].get("reporting_currency")),
                 ("Archetype", dcf["archetype"]["kind"]), ("Engine", dcf["archetype"].get("engine")),
                 ("Decisive variable", dcf["archetype"]["decisive"].get("label"))):
        w.label(k, v, fmt=NUM2 if isinstance(v, float) else None)
    w.r += 1
    hr = w.r
    for j, h in enumerate(["Scenario", "Probability", "WACC", "EV ($M)", "Equity ($M)", "Value / listed unit",
                           "Upside", "Model's stated value"]):
        w.put(hr, 1 + j, h, font=BOLD, fill=HEAD)
    w.r += 1
    r0 = w.r
    for scn in dcf["scenarios"]:
        name, cells = out[scn["id"]]
        w.put(w.r, 1, scn.get("label") or scn["id"])
        w.put(w.r, 2, scn.get("probability"), fmt=PCT, fill=YELLOW)
        w.put(w.r, 3, f"={_q(name, 'B5')}", fmt=PCT)
        w.put(w.r, 4, f"={_q(name, cells['ev'])}", fmt=NUM)
        w.put(w.r, 5, f"={_q(name, cells['equity'])}", fmt=NUM)
        w.put(w.r, 6, f"={_q(name, cells['vps'])}", fmt=NUM2, font=BOLD)
        w.put(w.r, 7, f"=F{w.r}/{dcf['market']['price']}-1", fmt=PCT)
        w.put(w.r, 8, (scn.get("results") or {}).get("value_per_share"), fmt=NUM2, font=MUTED)
        w.r += 1
    r1 = w.r - 1
    if all(s.get("probability") is not None for s in dcf["scenarios"]):
        w.put(w.r, 1, "Probability-weighted", font=BOLD)
        w.put(w.r, 6, f"=SUMPRODUCT(B{r0}:B{r1},F{r0}:F{r1})", fmt=NUM2, font=BOLD)
        w.r += 1
    w.r += 1
    w.put(w.r, 1, "Values recompute from each scenario sheet's yellow inputs. 'Model's stated value' is the "
                  "number the model itself reported (for imported models the engine may only approximate it).",
          font=MUTED)
    w.r += 2
    narr = dcf.get("narrative") or {}
    for k in ("headline", "decisive_variable", "thesis", "bear_case", "bull_case", "opposing_case", "market_implied"):
        if narr.get(k):
            w.put(w.r, 1, k.replace("_", " ").capitalize(), font=BOLD)
            c = w.put(w.r, 2, narr[k])
            c.alignment = Alignment(wrap_text=True, vertical="top")
            summary.merge_cells(start_row=w.r, start_column=2, end_row=w.r, end_column=8)
            summary.row_dimensions[w.r].height = min(15 * (1 + len(narr[k]) // 110), 160)
            w.r += 1
    summary.column_dimensions["A"].width = 28
    for col in "BCDEFGH":
        summary.column_dimensions[col].width = 16

    if dcf.get("checks"):
        ws = wb.create_sheet("Checks")
        for j, h in enumerate(["Check", "Status", "Detail", "Diagnosis"]):
            ws.cell(1, 1 + j, h).font = BOLD
        for i, c in enumerate(dcf["checks"], 2):
            ws.cell(i, 1, c.get("label") or c["id"])
            ws.cell(i, 2, c["status"])
            ws.cell(i, 3, c.get("detail"))
            ws.cell(i, 4, c.get("diagnosis"))
        ws.column_dimensions["A"].width = 44
        ws.column_dimensions["C"].width = 70
    if dcf.get("sources"):
        ws = wb.create_sheet("Sources")
        for j, h in enumerate(["Id", "Kind", "Description", "Accession", "URL"]):
            ws.cell(1, 1 + j, h).font = BOLD
        for i, s in enumerate(dcf["sources"], 2):
            for j, k in enumerate(("id", "kind", "description", "accession", "url")):
                ws.cell(i, 1 + j, s.get(k))
        ws.column_dimensions["C"].width = 50
        ws.column_dimensions["E"].width = 60
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def verify_export(dcf, *, tol=1e-6):
    """Recalculate the export through LibreOffice and compare per-scenario value
    per listed unit with dcfkit.engine. Returns {scenario: (sheet, engine, diff)}."""
    import copy
    import os
    import shutil
    import subprocess
    import tempfile

    from openpyxl import load_workbook

    from dcfkit import engine as E

    data = export(dcf)
    tmp = tempfile.mkdtemp(prefix="dcfx-")
    try:
        src = os.path.join(tmp, "in")
        out = os.path.join(tmp, "out")
        os.makedirs(src)
        os.makedirs(out)
        p = os.path.join(src, "model.xlsx")
        open(p, "wb").write(data)
        subprocess.run(["soffice", f"-env:UserInstallation=file://{tmp}/profile", "--headless", "--calc",
                        "--convert-to", "xlsx", "--outdir", out, p], capture_output=True, timeout=180, check=True)
        wb = load_workbook(os.path.join(out, "model.xlsx"), data_only=True)
        ref = E.value_model(copy.deepcopy(dcf))
        res = {}
        ws = wb["Summary"]
        rows = {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1)}
        for scn in ref["scenarios"]:
            r = rows.get(scn.get("label") or scn["id"])
            got = ws.cell(r, 6).value if r else None
            exp = scn["results"]["value_per_share"]
            res[scn["id"]] = (got, exp, None if got is None else abs(got - exp))
        return res
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
