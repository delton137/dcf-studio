"""Deterministic importer: legacy Grok DCF workbooks -> DCF Studio dcf.json.

    uv run python -m server.importer.grok_xlsx --all [--force] [--tickers AMKR ETN]

Source: DCF_models/grok_dcf_models/ (READ-ONLY -- files are copied, never moved,
and the build_*.py scripts are only read as text, never executed).

Every Grok workbook is a hardcoded-value export (zero formulas) with sheets
Summary / Assumptions / Historical / Proj_{Base,Bull,Bear} / Sensitivity /
Scenario_Compare / Notes_Methodology. Everything is located BY LABEL.

Mapping (one consolidated DCF component per scenario):
  * rows       <- Proj_<S> (segment revenues and other lines -> drivers{});
                  FCFF exactly as in the sheet, any identity gap -> other_adj
  * terminal   <- "Terminal method" text: normalized/plain Gordon with
                  fcf_next = TV x (WACC - g) (reproduces the sheet's TV exactly);
                  "Exit EV/EBITDA Nx" -> exit_multiple; a fixed residual floor
                  -> equivalent revenue multiple (documented in terminal.note)
  * discounting<- mid-year flows, TV at year N (confirmed per model with
                  engine.detect_convention)
  * bridge     <- the rows between EV and Equity Value (already signed equity
                  effects); net debt split into cash/debt when it foots exactly
  * results    <- the sheet's EV / equity / per-share (Summary headline values)
  * headline   <- probability-weighted value (weights from build_<t>.py text,
                  cross-checked against the Summary label and the PW value)
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import math
import re
import shutil
import sys
from datetime import date, datetime, timezone
from pathlib import Path

import openpyxl

from server import store
from server.store import LEGACY_ROOT, MODELS, DATA, now_iso, read_json, write_json, update_run
from dcfkit import engine as E
from dcfkit.schema import DCF

GROK_DIR = LEGACY_ROOT / "grok_dcf_models"
REPORT_PATH = DATA / "imports" / "grok_report.json"
SKIP = {"CRBS"}  # byte-level duplicate of CBRS (user ticker typo)
PSYCHEDELIC = {"ATAI", "CMPS", "DFTX", "GHRS", "HELP"}  # built by psychedelic_dcf_engine.py
GENERATION = "grok_xlsx"

# Listed-unit facts (ordinary shares per listed unit). Grok values these per
# ADS/ADR with an ADS-equivalent share count; we convert the count to ordinary
# shares (x ratio) so value_per_ordinary x ratio = the workbook's per-ADS value.
LISTING = {
    "TSM": ("ADR", 5.0),
    "SIMO": ("ADS", 4.0),
    "WRD": ("ADS", 3.0),
    "SNY": ("ADR", 0.5),
}

# Model period basis. Grok projects calendar years for most names (incl. the
# non-December NVDA/AVGO/MRVL/MU/IREN/APLD, as their sheets/notes state); these
# five project fiscal years. (MM-DD of fiscal year end, offset added to the
# sheet's year label to get the calendar year the fiscal year ends in.)
FISCAL = {
    "DELL": ("01-31", 0),   # FY2027 ends ~Jan 31, 2027
    "MSFT": ("06-30", 0),   # FY2027 ends Jun 30, 2027
    "LULU": ("01-31", 1),   # FY2026 ends ~Feb 1, 2027
    "COHR": ("06-30", 0),   # FY2026 ends ~Jun 30, 2026
    "SNDK": ("06-30", 0),   # FY2026 ended Jul 3, 2026 (52/53-wk year nearest Jun 30)
}
# Company fiscal year end where Grok nevertheless projects calendar years
COMPANY_FYE = {"NVDA": "01-31", "AVGO": "11-02", "MRVL": "01-31", "MU": "08-31",
               "IREN": "06-30", "APLD": "05-31"}
# Historical-table year basis for fiscal-year tables that carry FCF
HIST_FISCAL = {"NVDA": ("01-31", 0), "LULU": ("01-31", 1), "DELL": ("01-31", 0), "MSFT": ("06-30", 0),
               "COHR": ("06-30", 0), "SNDK": ("06-30", 0), "AVGO": ("11-02", 0)}

ARCHETYPE = {
    "AMBP": "manufacturer", "AMKR": "semiconductor", "APLD": "capital_intensive", "ASML": "semiconductor",
    "ATAI": "biotech_pipeline", "AVGO": "semiconductor", "BBAI": "software", "BE": "manufacturer",
    "CBRS": "semiconductor", "CEG": "utility", "CMPS": "biotech_pipeline", "COHR": "semiconductor",
    "CORZ": "capital_intensive", "CRWV": "capital_intensive", "DELL": "manufacturer",
    "DFTX": "biotech_pipeline", "ETN": "manufacturer", "GHRS": "biotech_pipeline", "GLW": "manufacturer",
    "HELP": "biotech_pipeline", "IREN": "capital_intensive", "LULU": "consumer", "MCD": "consumer",
    "META": "ad_platform", "MRNA": "biotech_pipeline", "MRVL": "semiconductor", "MSFT": "software",
    "MU": "semiconductor", "NBIS": "capital_intensive", "NVDA": "semiconductor", "ON": "semiconductor",
    "RGTI": "pre_revenue", "RIVN": "manufacturer", "RKLB": "manufacturer", "SIMO": "semiconductor",
    "SNDK": "semiconductor", "SNY": "other", "SPCX": "capital_intensive", "TSM": "semiconductor",
    "WM": "capital_intensive", "WRD": "other",
}

SCENARIOS = (("bear", "Bear"), ("base", "Base"), ("bull", "Bull"))
SUMMARY_COL = {"bear": 1, "base": 2, "bull": 3, "market": 4}


class ImportError_(RuntimeError):
    """The workbook cannot be mapped faithfully."""


# --------------------------------------------------------------------------
# small helpers
# --------------------------------------------------------------------------

def norm(s):
    s = str(s).replace("−", "-").replace("–", "-").replace("—", "-").replace("Δ", "d")
    s = s.replace("×", "x")
    return re.sub(r"\s+", " ", s).strip().lower()


def slug(s):
    return re.sub(r"[^a-z0-9]+", "_", norm(s)).strip("_")[:60] or "x"


def isnum(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(float(v))


def num(v):
    """Number, or parse '12.5%' / '9x' / '1,234' strings; else None."""
    if isnum(v):
        return float(v)
    if isinstance(v, str):
        s = v.strip().replace(",", "")
        m = re.fullmatch(r"([-+]?\d+(?:\.\d+)?)\s*(%|x|×)?", s)
        if m:
            x = float(m.group(1))
            return x / 100.0 if m.group(2) == "%" else x
    return None


def rel(p):
    return str(Path(p).resolve().relative_to(LEGACY_ROOT.resolve()))


def close(a, b, rtol=1e-6, atol=1e-6):
    return a is not None and b is not None and abs(a - b) <= max(atol, rtol * max(abs(a), abs(b)))


def money_scale(label):
    """Unit multiplier to millions from a label like 'Enterprise Value ($B)'."""
    n = norm(label)
    if re.search(r"\((\$|€|eur|usd)?\s*b\b|\$b\b|€b\b|\(b\)", n):
        return 1000.0
    return 1.0


MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep",
                                       "oct", "nov", "dec"], 1)}


def parse_date(text, default_year=None):
    m = re.search(r"\b(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s+(\d{1,2})(?:,?\s+(\d{4}))?",
                  str(text), re.I)
    if not m:
        return None
    year = int(m.group(3)) if m.group(3) else default_year
    if not year:
        return None
    try:
        return date(year, MONTHS[m.group(1).lower()[:3]], int(m.group(2)))
    except ValueError:
        return None


# --------------------------------------------------------------------------
# workbook access
# --------------------------------------------------------------------------

class Sheet:
    def __init__(self, ws):
        self.title = ws.title
        self.rows = [list(r) for r in ws.iter_rows(values_only=True)]

    def a(self, i):
        r = self.rows[i]
        return r[0] if r else None

    def cell(self, i, j):
        r = self.rows[i]
        return r[j] if r and j < len(r) else None

    def find(self, pattern, start=0, stop=None):
        rx = re.compile(pattern)
        for i in range(start, len(self.rows) if stop is None else min(stop, len(self.rows))):
            a = self.a(i)
            if a is not None and rx.search(norm(a)):
                return i
        return None

    def value(self, pattern, col=1, start=0, stop=None):
        i = self.find(pattern, start, stop)
        return (self.cell(i, col), i) if i is not None else (None, None)

    def texts(self):
        for r in self.rows:
            for v in r or []:
                if isinstance(v, str):
                    yield v


def load_book(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        sheets = {ws.title: Sheet(ws) for ws in wb.worksheets}
    finally:
        wb.close()
    return sheets


def hist_sheet(book):
    for name in book:
        if name.lower().startswith("historical"):
            return book[name]
    return None


# --------------------------------------------------------------------------
# inventory
# --------------------------------------------------------------------------

def inventory():
    """Legacy Grok models: [{ticker, xlsx, readme, csv, build_py, mtime}]."""
    items = []
    for x in sorted(GROK_DIR.glob("*_DCF_Model.xlsx")):
        t = x.name.split("_DCF_Model")[0].upper()
        if t in SKIP:
            continue
        readme = GROK_DIR / f"{t}_DCF_README.md"
        csv = GROK_DIR / f"{t}_DCF_Base_FCFF.csv"
        build = GROK_DIR / f"build_{t.lower()}_dcf.py"
        items.append({
            "ticker": t, "xlsx": x,
            "readme": readme if readme.exists() else None,
            "csv": csv if csv.exists() else None,
            "build_py": build if build.exists() else None,
            "mtime": datetime.fromtimestamp(x.stat().st_mtime),
        })
    return items


def find_item(ticker):
    for it in inventory():
        if it["ticker"] == ticker.upper():
            return it
    return None


# --------------------------------------------------------------------------
# Summary sheet
# --------------------------------------------------------------------------

def parse_summary(sh):
    hdr = None
    for i, r in enumerate(sh.rows):
        if r and len(r) > 3 and r[1] == "Bear" and r[2] == "Base" and r[3] == "Bull":
            hdr = i
            break
    if hdr is None:
        raise ImportError_("Summary: Bear/Base/Bull header row not found")
    out = {"hdr": hdr, "snapshot": [], "block": [], "pw": {}, "footnotes": []}
    # market snapshot: rows before the first big section after row 4
    for i in range(3, hdr):
        a = sh.a(i)
        b = sh.cell(i, 1)
        if a is None:
            continue
        out["snapshot"].append((str(a).strip(), b, i))
    pw_i = sh.find(r"^probability-weighted fair value$", hdr) or sh.find(r"probability-weighted", hdr)
    stop = pw_i if pw_i is not None else len(sh.rows)
    for i in range(hdr + 1, stop):
        a = sh.a(i)
        if a is None:
            continue
        vals = [sh.cell(i, c) for c in (1, 2, 3, 4)]
        out["block"].append((str(a).strip(), vals, i))
    verdict_i = sh.find(r"^verdict$", hdr)
    if pw_i is not None:
        for i in range(pw_i + 1, verdict_i if verdict_i else len(sh.rows)):
            a = sh.a(i)
            if a is None:
                continue
            la = norm(a)
            v = sh.cell(i, 1)
            if la.startswith("*"):
                out["footnotes"].append((str(a).strip(), v))
            elif "weights" in la:
                out["pw"]["weights_text"] = v
            elif re.match(r"^(bear|base|bull) weight$", la):
                out["pw"].setdefault("weights_rows", {})[la.split()[0]] = num(v)
            elif la.startswith("probability-weighted fv") or la.startswith("probability-weighted fair value /"):
                out["pw"]["fv"] = num(v)
                out["pw"]["fv_label"] = str(a).strip()
            elif "equity value" in la:
                out["pw"]["equity"] = num(v) * money_scale(a) if num(v) is not None else None
            elif la.startswith("vs ") or la.startswith("implied market"):
                out["pw"].setdefault("vs", []).append((str(a).strip(), v))
    if verdict_i is not None:
        out["verdict"] = str(sh.cell(verdict_i, 1) or "").strip()
        for i in range(verdict_i + 1, min(verdict_i + 6, len(sh.rows))):
            if sh.a(i):
                out["conclusion"] = str(sh.a(i)).strip()
                break
    kd = sh.find(r"^key debates?\b", hdr)
    if kd is not None:
        txt = str(sh.a(kd))
        if norm(txt) in ("key debate", "key debates"):
            for i in range(kd + 1, min(kd + 4, len(sh.rows))):
                if sh.a(i):
                    txt = str(sh.a(i))
                    break
        out["key_debates"] = txt.strip()
    return out


def block_row(summ, pattern, *, exclude=None):
    rx = re.compile(pattern)
    for lab, vals, i in summ["block"]:
        n = norm(lab)
        if rx.search(n) and not (exclude and re.search(exclude, n)):
            return lab, vals
    return None, None


def snapshot_value(summ, pattern):
    rx = re.compile(pattern)
    for lab, v, i in summ["snapshot"]:
        if rx.search(norm(lab)) and num(v) is not None:
            return lab, num(v)
    return None, None


# --------------------------------------------------------------------------
# Proj_<Scenario> sheets
# --------------------------------------------------------------------------

CORE = {
    "revenue": r"^(total revenue|total net revenue|total company revenue|total net sales|revenue)$",
    "ebitda": r"^(adj\. ebitda|ebitda|total adj\. ebitda)$",
    "da": r"^\((-|\+)\) d&a$",
    "ebit": r"^ebit$",
    "taxes": r"^\(-\) cash taxes",
    "nopat": r"^nopat$",
    "capex": r"^\(-\) (total )?capex$",
    "dnwc": r"^\(-\) d ?(nwc|net working capital)$",
    "fcff": r"^unlevered fcff$",
    "df": r"^discount factor",
    "pv": r"^pv of fcff$",
}
RATIO_WORDS = r"%|margin|growth|turn|ratio|factor|\(units\)|asp|\(\$k\)"


def parse_proj(sh):
    hdr = None
    for i, r in enumerate(sh.rows[:12]):
        if r and isinstance(r[0], str) and norm(r[0]) in ("year", "calendar year", "fiscal year") and isnum(r[1]):
            hdr = i
            break
    if hdr is None:
        raise ImportError_(f"{sh.title}: year header row not found")
    years = []
    for v in sh.rows[hdr][1:]:
        if isnum(v):
            years.append(int(v))
        else:
            break
    n = len(years)
    lines, section = [], None
    vb = sh.find(r"^valuation bridge$", hdr)
    if vb is None:
        raise ImportError_(f"{sh.title}: VALUATION BRIDGE not found")
    for i in range(hdr + 1, vb):
        r = sh.rows[i]
        if not r or r[0] is None:
            continue
        vals = r[1:1 + n]
        nums = [v for v in vals if isnum(v)]
        if not nums and all(v is None for v in vals):
            section = str(r[0]).strip()
            continue
        if len(nums) == n:
            lines.append({"label": str(r[0]).strip(), "values": [float(v) for v in vals],
                          "section": section, "row": i + 1})
    bridge = []
    for i in range(vb + 1, len(sh.rows)):
        r = sh.rows[i]
        if not r or r[0] is None:
            continue
        nums = [v for v in r[1:] if isnum(v)]
        if len(nums) > 1:      # trailing per-year rows (psychedelic engine) -- not part of the bridge
            continue
        bridge.append((str(r[0]).strip(), r[1] if len(r) > 1 else None, i + 1))
    scale = 1000.0 if any(norm(l["section"] or "").startswith("fcff bridge") and money_scale(l["section"]) > 1
                          for l in lines) else 1.0
    return {"title": sh.title, "years": years, "lines": lines, "bridge": bridge, "scale": scale,
            "head": str(sh.a(0) or ""), "narrative": str(sh.a(1) or "") if len(sh.rows) > 1 else ""}


def proj_core(p):
    core, used = {}, set()
    for key, pat in CORE.items():
        rx = re.compile(pat)
        for k, l in enumerate(p["lines"]):
            if k in used:
                continue
            if rx.search(norm(l["label"])):
                if key == "revenue" and norm(l["label"]) == "revenue":
                    # prefer an explicit "Total ..." line when present
                    tot = [j for j, x in enumerate(p["lines"]) if norm(x["label"]).startswith("total")
                           and re.search(r"revenue|sales", norm(x["label"]))]
                    if tot:
                        k, l = tot[0], p["lines"][tot[0]]
                core[key] = (k, l)
                used.add(k)
                break
    # second D&A row ("(+) D&A") is a duplicate of "(-) D&A"
    for k, l in enumerate(p["lines"]):
        if k not in used and re.search(CORE["da"], norm(l["label"])):
            used.add(k)
    missing = [k for k in ("revenue", "da", "ebit", "taxes", "nopat", "capex", "dnwc", "fcff") if k not in core]
    if missing:
        raise ImportError_(f"{p['title']}: rows not found: {missing}")
    return core, used


def proj_bridge(p):
    """Valuation-bridge block -> dict (millions)."""
    sc = p["scale"]
    out = {"items": [], "fv_rows": [], "other_rows": []}
    ev_i = eq_i = None
    for k, (lab, v, row) in enumerate(p["bridge"]):
        n = norm(lab)
        x = num(v)
        if n.startswith("sum of pv"):
            out["pv_explicit"] = x * money_scale(lab)
        elif n.startswith("normalized terminal fcf"):
            out["norm_fcf"] = x * money_scale(lab)
        elif n.startswith("terminal value"):
            out["tv"] = x * money_scale(lab)
        elif n == "terminal method":
            out["tv_method"] = str(v)
        elif n.startswith("pv of terminal"):
            out["pv_tv"] = x * money_scale(lab)
        elif n.startswith("enterprise value"):
            # prefer the ($M) row when both ($M) and ($B) are shown (TSM/AVGO/NVDA)
            if "ev" not in out or (money_scale(lab) == 1.0 and money_scale(out["ev_label"]) > 1.0):
                out["ev"] = x * money_scale(lab)
                out["ev_label"] = lab
            ev_i = k if ev_i is None else ev_i
        elif n.startswith("equity value"):
            if eq_i is None:
                out["equity"] = x * money_scale(lab)
                out["equity_label"] = lab
                eq_i = k
        elif ev_i is not None and eq_i is None and (n.startswith("(-)") or n.startswith("(+)")):
            out["items"].append((lab, x * money_scale(lab)))
        elif n.startswith("diluted"):
            out["shares"] = x * (1000.0 if re.search(r"\(b\)", n) else 1.0)
            out["shares_label"] = lab
        elif re.match(r"^(fair value per|fair value /)", n):
            out.setdefault("fv", x)
            out.setdefault("fv_label", lab)
        elif n.startswith("current") and "price" in n:
            out["price"] = x
        elif n == "wacc":
            out["wacc"] = x
        elif n == "terminal g":
            out["g"] = x
        elif n.startswith("% of ev"):
            out["tv_share"] = x
        elif n.startswith("upside"):
            out["upside"] = x
        elif x is not None:
            out["other_rows"].append((lab, x))
    for key in ("pv_explicit", "tv", "pv_tv", "ev", "equity", "wacc"):
        if out.get(key) is None:
            raise ImportError_(f"{p['title']}: bridge value {key!r} not found")
    return out


def g_from_method_text(method_text):
    m = re.search(r"g\s*=\s*([-\d.]+)%", method_text or "")
    return float(m.group(1)) / 100.0 if m else None


# --------------------------------------------------------------------------
# build-script / README text
# --------------------------------------------------------------------------

def weights_from_build(item):
    """(bear, base, bull) from build_<t>.py text (or CFG['pw_weights']); never executed."""
    p = item.get("build_py")
    if not p:
        return None, None
    txt = Path(p).read_text(errors="replace")
    m = re.search(r"w_bear\s*,\s*w_base\s*,\s*w_bull\s*=\s*\(?\s*([\d.]+)\s*,\s*([\d.]+)\s*,\s*([\d.]+)", txt)
    if not m:
        m = re.search(r"[\"']pw_weights[\"']\s*:\s*\(\s*([\d.]+)\s*,\s*([\d.]+)\s*,\s*([\d.]+)", txt)
    if not m:
        return None, None
    return tuple(float(x) for x in m.groups()), rel(p)


def hidden_lease(item):
    """OPLEASE constant from the build script text (e.g. 'OPLEASE = 0.5 + 2.3'); never executed."""
    p = item.get("build_py")
    if not p:
        return None
    m = re.search(r"^OPLEASE\s*=\s*([\d_.+\s]+?)\s*(#.*)?$", Path(p).read_text(errors="replace"), re.M)
    if not m:
        return None
    try:
        return sum(float(x.replace("_", "")) for x in m.group(1).split("+") if x.strip())
    except ValueError:
        return None


def weights_from_text(s):
    if not s:
        return None
    m = re.search(r"(\d+(?:\.\d+)?)%?\s*/\s*(\d+(?:\.\d+)?)%?\s*/\s*(\d+(?:\.\d+)?)%?", str(s))
    if not m:
        return None
    w = [float(x) for x in m.groups()]
    if sum(w) > 1.5:
        w = [x / 100.0 for x in w]
    return tuple(w)


def read_readme(item):
    p = item.get("readme")
    return Path(p).read_text(errors="replace") if p else ""


def readme_field(txt, name):
    m = re.search(rf"\*\*{name}:\*\*\s*(.+)", txt)
    return m.group(1).strip() if m else None


def notes_sections(sh):
    """Notes_Methodology -> {SECTION TITLE: text}."""
    secs, cur = {}, "_"
    if sh is None:
        return secs
    for r in sh.rows[1:]:
        a = r[0] if r else None
        if not a:
            continue
        s = str(a)
        m = re.match(r"^\s*\d+\.\s+([A-Z][A-Z0-9 &/'()\-—,.]+.*)$", s)
        m2 = re.match(r"^([A-Z][A-Z &/'-]{3,}):\s*(.*)$", s)
        if m and not s.startswith(" "):
            cur = m.group(1).strip()
            secs[cur] = ""
        elif m2:
            cur = m2.group(1).strip()
            secs[cur] = m2.group(2).strip()
        else:
            secs[cur] = (secs.get(cur, "") + " " + s.strip()).strip()
    return secs


# --------------------------------------------------------------------------
# conversion
# --------------------------------------------------------------------------

def period_for(ticker, year):
    if ticker in FISCAL:
        mmdd, off = FISCAL[ticker]
        y = year + off
        return f"FY{year}", date(y, int(mmdd[:2]), int(mmdd[3:]))
    return str(year), date(year, 12, 31)


def driver_meta(label, section, unit_ccy, scale):
    n = norm(label)
    if re.search(r"%|margin|growth", n):
        return "ratio", "margin" if "margin" in n else "ratio"
    if "turn" in n:
        return "x", "capital"
    if "units" in n or "deliveries" in n:
        return "units", "volume"
    if "asp" in n:
        return f"{unit_ccy}k" if "$k" in n else unit_ccy, "price"
    if "pp&e" in n or "ppe" in n or "capex" in n:
        return f"{unit_ccy}M", "capital"
    return f"{unit_ccy}M", "other"


def convert(item):
    """Convert one inventory item to a dcf.json dict (validated by DCF.model_validate)."""
    t = item["ticker"]
    book = load_book(item["xlsx"])
    for need in ("Summary", "Assumptions", "Proj_Base", "Proj_Bull", "Proj_Bear"):
        if need not in book:
            raise ImportError_(f"sheet {need!r} missing")
    S, A = book["Summary"], book["Assumptions"]
    H = hist_sheet(book)
    notes = notes_sections(book.get("Notes_Methodology"))
    readme = read_readme(item)
    summ = parse_summary(S)
    import_notes = []

    # ---------------- header / dates / company
    a2 = str(S.a(1) or "")
    parts = [x.strip() for x in a2.split("|")]
    m = re.match(r"^(.*?)\s*\((nasdaq|nyse|nyse american)\b[:\s]*([^)]*)\)", parts[0], re.I)
    company = (m.group(1) if m else parts[0]).strip() or t
    exchange = None
    if m:
        exchange = {"nasdaq": "Nasdaq", "nyse": "NYSE", "nyse american": "NYSE American"}[m.group(2).lower()]
    as_of = None
    for src in [p for p in parts if p.lower().startswith("as of")] + [readme_field(readme, "As of") or ""]:
        as_of = parse_date(src)
        if as_of:
            break
    mtime = item["mtime"]
    if as_of is None:
        as_of = mtime.date()
        import_notes.append("as_of from file mtime (no 'As of' header found)")
    latest = next((p.split(":", 1)[1].strip() for p in parts if p.lower().startswith("latest")), None) \
        or readme_field(readme, "Latest") or next((p for p in parts[2:3]), None) or "n/a"
    latest = re.sub(r"\*\*", "", latest)

    # ---------------- currencies / listing
    labels_all = " ".join(norm(x) for x in S.texts())
    rep_ccy = "EUR" if "€m" in labels_all else "USD"
    unit_ccy = "€" if rep_ccy == "EUR" else "$"
    listed_unit, ratio = LISTING.get(t, ("common", 1.0))
    fx = 1.0
    if rep_ccy != "USD":
        _, fxv = snapshot_value(summ, r"^eurusd")
        if fxv is None:
            v, _ = A.value(r"^eurusd")
            fxv = num(v)
        if fxv is None:
            raise ImportError_("non-USD model without an FX rate")
        fx = fxv

    # ---------------- scenarios from Proj sheets
    projs = {sid: parse_proj(book[f"Proj_{lab}"]) for sid, lab in SCENARIOS}
    years = projs["base"]["years"]
    for sid in ("bear", "bull"):
        if projs[sid]["years"] != years:
            raise ImportError_(f"Proj_{sid} years differ from Proj_Base")
    scale = projs["base"]["scale"]
    if scale != 1.0:
        import_notes.append(f"workbook flows are in $B; scaled x{scale:g} to millions")

    # Summary scenario block
    fv_lab, fv_vals = block_row(summ, r"^fair value /")
    if fv_lab is None:
        raise ImportError_("Summary: 'Fair Value / ...' row not found")
    ev_lab, ev_vals = block_row(summ, r"^enterprise value")
    eq_lab, eq_vals = block_row(summ, r"^equity value")
    sh_lab, sh_vals = block_row(summ, r"^diluted")
    up_lab, up_vals = block_row(summ, r"^upside")
    w_lab, w_vals = block_row(summ, r"^wacc$")
    g_lab, g_vals = block_row(summ, r"^terminal (g|growth)")

    weights, wsrc = weights_from_build(item)
    wtxt = weights_from_text(summ["pw"].get("weights_text"))
    if summ["pw"].get("weights_rows"):
        wr = summ["pw"]["weights_rows"]
        wtxt = (wr.get("bear"), wr.get("base"), wr.get("bull"))
    if weights is None and wtxt is not None:
        weights, wsrc = wtxt, "Summary label"
    if weights is not None and wtxt is not None and any(abs(a - b) > 1e-9 for a, b in zip(weights, wtxt)):
        import_notes.append(f"PW weights in {wsrc} {weights} differ from Summary label {wtxt}; using {wsrc}")
    pw_legacy = summ["pw"].get("fv")

    scenarios, bridge_by_scn, results_by_scn, rows_by_scn = [], {}, {}, {}
    drivers_def = {}
    shares_listed_by_scn = {}
    for sid, lab in SCENARIOS:
        p = projs[sid]
        core, used = proj_core(p)
        br = proj_bridge(p)
        sc = p["scale"]
        col = SUMMARY_COL[sid] - 1
        vps = num(fv_vals[col])
        if vps is None:
            raise ImportError_(f"Summary {fv_lab!r}: no {sid} value")
        ev_s = num(ev_vals[col]) * money_scale(ev_lab) if ev_lab else None
        eq_s = num(eq_vals[col]) * money_scale(eq_lab) if eq_lab else None
        if ev_s is not None and not close(ev_s, br["ev"], 1e-9, 1e-6):
            import_notes.append(f"{sid}: Summary EV {ev_s} != Proj EV {br['ev']}; using Proj")
        if eq_s is not None and not close(eq_s, br["equity"], 1e-9, 1e-6):
            import_notes.append(f"{sid}: Summary equity {eq_s} != Proj equity {br['equity']}; using Proj")
        if br.get("fv") is not None and not close(br["fv"], vps, 1e-9, 1e-9):
            import_notes.append(f"{sid}: Proj per-share {br['fv']} != Summary {vps}; headline uses Summary")
        wacc = br["wacc"]
        tv_text = br.get("tv_method") or ""
        g = br.get("g")
        if g is None and g_lab:
            g = num(g_vals[col])
        if g is None:
            g = g_from_method_text(tv_text)

        # ---- rows
        rev_k, rev_l = core["revenue"]
        seg_rows, seg_keys = [], []
        rev_section = rev_l["section"]
        for k, l in enumerate(p["lines"]):
            if k >= rev_k or k in used:
                continue
            if l["section"] == rev_section and rev_section and re.search(r"revenue|sales", norm(rev_section)) \
                    and not re.search(RATIO_WORDS, norm(l["label"])):
                seg_rows.append(k)
        if seg_rows:
            sums = [sum(p["lines"][k]["values"][j] for k in seg_rows) for j in range(len(years))]
            if not all(close(a, b, 1e-6, 1e-3) for a, b in zip(sums, rev_l["values"])):
                seg_rows = []   # not a clean segment split -> plain drivers
        rows = []
        other_adj_years = []
        for j, y in enumerate(years):
            def v(key):
                return core[key][1]["values"][j] * sc if key in core else None
            revenue, ebit, da = v("revenue"), v("ebit"), v("da")
            taxes, capex, dnwc, fcff = v("taxes"), v("capex"), v("dnwc"), v("fcff")
            nopat_sheet = v("nopat")
            ebitda = v("ebitda") if "ebitda" in core else ebit + da
            other = 0.0
            if taxes < 0:
                other += -taxes        # tax credit -> other_adj, keeps taxes >= 0
                taxes = 0.0
            nopat = ebit - taxes
            if nopat_sheet is not None and not close(nopat_sheet, ebit - v("taxes"), 1e-9, 1e-6):
                import_notes.append(f"{sid} {y}: sheet NOPAT {nopat_sheet} != EBIT - taxes")
            ident = nopat + da - capex - dnwc + other
            other += fcff - ident
            if abs(other) < 1e-9 * max(1.0, abs(fcff)):
                other = 0.0
            if other:
                other_adj_years.append((y, other))
            drivers = {}
            for k, l in enumerate(p["lines"]):
                if k in used:
                    continue
                val = l["values"][j]
                if k in seg_rows:
                    key = f"seg_rev:{l['label']}"
                    drivers[key] = val * sc
                    drivers_def.setdefault(key, {"key": key, "label": l["label"], "unit": f"{unit_ccy}M",
                                                 "kind": "segment"})
                else:
                    sect = norm(l["section"] or "")
                    prefix = ""
                    if "ebitda" in sect and "ebitda" in norm(l["label"]) \
                            and not re.search(RATIO_WORDS, norm(l["label"])):
                        prefix = "seg_ebitda:"
                    key = prefix + slug(l["label"]) if not prefix else prefix + l["label"]
                    is_ratio = bool(re.search(RATIO_WORDS, norm(l["label"])))
                    drivers[key] = val if is_ratio else val * sc
                    unit, kind = driver_meta(l["label"], l["section"], unit_ccy, sc)
                    drivers_def.setdefault(key, {"key": key, "label": l["label"], "unit": unit, "kind": kind})
            period, pend = period_for(t, y)
            rows.append({
                "period": period, "period_end": pend.isoformat(), "is_actual": False,
                "revenue": revenue, "ebitda": ebitda, "ebit": ebit, "taxes": taxes, "nopat": nopat,
                "da": da, "capex": capex, "delta_nwc": dnwc, "other_adj": other, "fcff": fcff,
                "drivers": drivers,
            })
        if other_adj_years:
            mx = max(abs(x) for _, x in other_adj_years)
            if mx > 0.01:
                import_notes.append(f"{sid}: sheet FCFF differs from NOPAT+D&A-capex-dNWC in "
                                    f"{len(other_adj_years)} year(s) (max {mx:.3f}M) -> other_adj")
        rows_by_scn[sid] = rows

        # ---- terminal
        tv, pv_tv = br["tv"], br["pv_tv"]
        last = rows[-1]
        mt = norm(tv_text)
        term = None
        mm = re.search(r"exit ev/ebitda\s*([\d.]+)x", mt)
        if mm and last["ebitda"] and close(float(mm.group(1)) * max(last["ebitda"], 0.0), tv, 1e-9, 1e-6):
            term = {"method": "exit_multiple", "multiple": float(mm.group(1)), "multiple_metric": "ebitda",
                    "metric_period": "last", "g": g,
                    "note": f"legacy: {tv_text}"}
        elif "gordon" in mt and g is not None:
            meth = "normalized_gordon" if "normalized" in mt else "gordon"
            term = {"method": meth, "g": g, "fcf_next": tv * (wacc - g),
                    "note": f"legacy: {tv_text}; fcf_next = TV x (WACC - g) reproduces the sheet TV"
                            + (f" (sheet 'Normalized terminal FCF' {br['norm_fcf']:.2f} is the year-N "
                               f"basis; TV = it x (1+g)/(WACC-g))" if br.get("norm_fcf") is not None else "")}
        elif last["revenue"] and last["revenue"] > 0:
            term = {"method": "exit_multiple", "multiple": tv / last["revenue"], "multiple_metric": "revenue",
                    "metric_period": "last", "g": g,
                    "note": f"legacy: {tv_text} -- a fixed terminal amount of {tv:.1f}M (Grok build: "
                            f"max(20% x cash, $50M) when terminal EBITDA <= 0); encoded as the equivalent "
                            f"revenue multiple so the engine reproduces it and WACC/g knobs leave it fixed"}
            import_notes.append(f"{sid}: terminal '{tv_text}' (fixed TV {tv:.1f}M) encoded as revenue multiple")
        else:
            raise ImportError_(f"{sid}: cannot map terminal method {tv_text!r}")
        if term.get("g") is None:
            term.pop("g")

        # ---- shares (sheet shares are per LISTED unit -> ordinary)
        shares_listed = br.get("shares")
        if shares_listed is None and sh_lab:
            shares_listed = num(sh_vals[col]) * (1000.0 if re.search(r"\(b\)", norm(sh_lab)) else 1.0)
        if shares_listed is None:
            # SNY: FV/ADR = equity x fx / ADR count
            shares_listed = br["equity"] * fx / vps
            import_notes.append(f"share count derived as equity x fx / per-listed-unit value "
                                f"({shares_listed:.2f}M listed units; no share row in Proj sheets)")
        shares_listed_by_scn[sid] = shares_listed
        shares_ord = shares_listed * ratio
        eng_vps = br["equity"] / shares_ord * ratio * fx
        if not close(eng_vps, vps, 1e-6, 1e-6):
            import_notes.append(f"{sid}: equity/shares x ratio x fx = {eng_vps:.4f} vs sheet {vps:.4f}")

        # ---- bridge items (signed equity effects); must foot EV -> equity
        items = br["items"]
        tot = sum(x for _, x in items)
        if not close(br["ev"] + tot, br["equity"], 1e-9, 1e-6):
            flipped = [(l, -x if norm(l).startswith("(-)") and x > 0 else x) for l, x in items]
            if close(br["ev"] + sum(x for _, x in flipped), br["equity"], 1e-9, 1e-6):
                items = flipped
            else:
                resid = br["equity"] - br["ev"] - tot
                hidden = hidden_lease(item)
                if hidden is not None and close(-hidden, resid, 1e-9, 1e-6):
                    lab_h = "(−) Operating leases (build script OPLEASE; not shown in workbook) ($M)"
                    import_notes.append(f"bridge: workbook equity includes {-resid:.2f}M operating leases that "
                                        f"only the build script shows (OPLEASE); added as a bridge item")
                elif abs(resid) <= 0.01 * max(abs(br["equity"]), abs(br["ev"]), 1.0):
                    lab_h = "Unlabelled bridge residual (workbook EV -> equity) ($M)"
                    import_notes.append(f"bridge: unlabelled residual {resid:.4f}M between EV and equity "
                                        f"added as an 'other' item")
                else:
                    raise ImportError_(f"{sid}: EV {br['ev']} + bridge {tot} != equity {br['equity']}")
                items = items + [(lab_h, resid)]
        bridge_by_scn[sid] = items

        disc = {"wacc": wacc, "convention": "mid", "stub_fraction": 1.0, "tv_timing": "end_last"}
        up = num(up_vals[col]) if up_lab else None
        res = {
            "ev": br["ev"], "equity": br["equity"], "value_per_share": vps, "value_per_share_raw": vps,
            "value_per_ordinary": br["equity"] / shares_ord, "shares_used": shares_ord,
            "tv_share_of_ev": (pv_tv / br["ev"]) if br["ev"] > 0 else None,
            "upside": up if up is not None else vps / num(S.cell(4, 1)) - 1.0,
            "pv_explicit": br["pv_explicit"], "pv_tv": pv_tv, "tv": tv,
            "bridge_total": br["equity"] - br["ev"], "floored": False,
        }
        if t == "SNY":
            fvo = block_row(summ, r"^fv / ordinary")[1]
            if fvo:
                res["legacy_fv_per_ordinary_eur"] = num(fvo[col])
        results_by_scn[sid] = res
        comp = {
            "id": "consolidated", "label": "Consolidated FCFF (Grok legacy)", "kind": "dcf",
            "rows": rows, "terminal": term,
            "results": {"pv_explicit": br["pv_explicit"], "tv": tv, "pv_tv": pv_tv,
                        "value": br["pv_explicit"] + pv_tv},
        }
        narrative = p["narrative"].strip() or None
        scenarios.append({
            "id": sid, "role": sid, "label": lab,
            "probability": weights[[s for s, _ in SCENARIOS].index(sid)] if weights else None,
            "narrative": narrative, "year1_locked": False, "discounting": disc,
            "components": [comp],
            "shares": {"used": shares_ord, "method": "current_diluted",
                       "note": f"sheet '{br.get('shares_label') or sh_lab or 'derived'}' = {shares_listed:g}M"
                               + (f" {listed_unit}s x {ratio:g} ordinary per {listed_unit}" if ratio != 1 else "")},
            "results": res,
        })

    # ---------------- bridge (top level)
    bridge, overrides = build_bridge(t, summ, A, bridge_by_scn, import_notes)
    for s in scenarios:
        s["bridge_overrides"] = overrides.get(s["id"], {})

    # ---------------- market
    price = num(S.cell(4, 1))
    price_lab = str(S.a(4) or "")
    if price is None or not re.search(r"price", norm(price_lab)):
        raise ImportError_(f"Summary B5 is not a price: {price_lab!r}")
    pdt = parse_date(price_lab, as_of.year) or as_of
    base_res = results_by_scn["base"]
    shares_diluted = base_res["shares_used"]
    market = {"price": price, "price_ts": pdt.isoformat(),
              "price_source": f"Grok workbook Summary!B5 '{price_lab}'",
              "session": "close" if "close" in norm(price_lab) else None,
              "fx_spot": fx, "shares_diluted": shares_diluted, "extra": {}}
    if fx != 1.0:
        market["fx_ts"] = as_of.isoformat()
    mkt_shares = num(sh_vals[3]) if sh_lab else None
    if mkt_shares is not None:
        mkt_shares *= (1000.0 if re.search(r"\(b\)", norm(sh_lab)) else 1.0) * ratio
        market["shares_basic"] = mkt_shares
    mc_lab, mc = snapshot_value(summ, r"^market cap")
    if mc is not None:
        mc_m = mc * (1000.0 if re.search(r"\(\$?b\)|\$b|\(b\)", norm(mc_lab)) else 1.0)
        implied = price * shares_diluted / ratio
        if "€" in mc_lab:
            market["extra"]["legacy_market_cap_eur_m"] = mc_m
        elif abs(implied / mc_m - 1) <= 0.05:
            market["market_cap"] = mc_m
        else:
            market["extra"]["legacy_market_cap_m"] = mc_m
            market["extra"]["legacy_market_cap_label"] = mc_lab
    if t == "SNY":
        _, lp = snapshot_value(summ, r"ordinary price")
        if lp:
            market["local_price"] = lp
        v, _ = A.value(r"^ordinary diluted")
        if num(v):
            market["extra"]["ordinary_diluted_m"] = num(v)
    if market["session"] is None:
        market.pop("session")
    if not market["extra"]:
        market.pop("extra")

    # ---------------- basis
    margin_basis = "adjusted" if any(norm(core_label(projs["base"], "ebitda") or "").startswith(x)
                                     for x in ("adj", "total adj")) else "gaap"
    meth_text = " ".join(str(x) for x in A.texts() if "ethodology" in str(x)) + " " + " ".join(notes.values())
    sbc = "unknown"
    if re.search(r"sbc[^.]*not added back|sbc (is )?treated as (an )?economic cost|not added back", meth_text, re.I):
        sbc = "in_ebit"
    fye = FISCAL[t][0] if t in FISCAL else "12-31"
    basis = {"as_of": as_of.isoformat(), "valuation_date": as_of.isoformat(), "latest_period": latest,
             "reporting_currency": rep_ccy, "fiscal_year_end": fye, "margin_basis": margin_basis,
             "sbc_treatment": sbc, "extra": {}}
    if t in COMPANY_FYE:
        basis["extra"]["company_fiscal_year_end"] = COMPANY_FYE[t]
        basis["extra"]["model_periods"] = "calendar years (as projected by Grok)"
    if t in FISCAL:
        basis["extra"]["model_periods"] = "fiscal years"
    if not basis["extra"]:
        basis.pop("extra")

    # ---------------- archetype
    seg_labels = [d["label"] for d in drivers_def.values() if d["kind"] == "segment"]
    mtxt = " ".join(str(x) for x in A.texts())
    mrev = re.search(r"Revenue\s*=\s*([^.]+)", mtxt)
    rev_identity = ("Revenue = " + mrev.group(1).strip()) if mrev else \
        ("Revenue = " + " + ".join(seg_labels) if seg_labels else "Revenue (single line)")
    engine_desc = f"{rev_identity}; FCFF = EBIT(1-t) + D&A - capex - dNWC (Grok)"
    decisive = decisive_label(mtxt, book, summ)
    kind = ARCHETYPE.get(t)
    if kind is None:
        low = (mtxt + " " + readme).lower()
        kind = "biotech_pipeline" if "phase 3" in low else "semiconductor" if "semiconductor" in low else "other"

    # ---------------- headline
    pw_engine_style = sum(w * results_by_scn[s]["value_per_share"] for w, (s, _) in zip(weights, SCENARIOS)) \
        if weights else None
    headline = {"method": "base_dcf", "rationale": "", "value_per_share": base_res["value_per_share"]}
    if weights and pw_legacy is not None:
        if not close(pw_engine_style, pw_legacy, 1e-6, 1e-6):
            import_notes.append(f"PW check: sum(w x FV) = {pw_engine_style:.4f} vs Summary {pw_legacy:.4f}")
        headline = {"method": "prob_weighted",
                    "rationale": f"Grok probability-weighted fair value, weights bear/base/bull "
                                 f"{'/'.join(f'{w:.0%}' for w in weights)} (from {wsrc}). Legacy PW uses the "
                                 f"unfloored scenario values.",
                    "value_per_share": pw_legacy, "probability_weighted": pw_legacy}
        if any(r["value_per_share"] < 0 for r in results_by_scn.values()):
            import_notes.append("a scenario has negative equity value per share; legacy values kept unfloored "
                                "(engine/derived PW floors at 0 and will differ)")

    # ---------------- alt valuations, grids, historicals, narrative, sources, assumptions
    alts = alt_valuations(summ, fv_lab, A, unit_ccy)
    grids = parse_grids(book.get("Sensitivity"), base_res["value_per_share"], scenarios[1]["discounting"]["wacc"],
                        scenarios[1]["components"][0]["terminal"].get("g"), import_notes)
    hist = parse_historicals(t, H, unit_ccy)
    if mkt_shares:
        hist.setdefault("share_count_series", []).append(
            {"date": as_of.isoformat(), "shares": mkt_shares, "kind": "basic (workbook Market column)"})
    narrative = build_narrative(t, summ, readme, A, notes, pw_legacy, price, base_res["value_per_share"], decisive,
                                unit_ccy if listed_unit == "common" or rep_ccy == "USD" else "$")
    sources = build_sources(notes, readme)
    assumptions = build_assumptions(A)
    wacc_build = wacc_build_from(A)
    if wacc_build:
        scenarios[1]["discounting"]["wacc_build"] = wacc_build

    # ---------------- assemble
    legacy_created = datetime.fromtimestamp(Path(item["xlsx"]).stat().st_mtime, tz=timezone.utc).astimezone()
    files = [rel(item["xlsx"])] + [rel(item[k]) for k in ("readme", "csv", "build_py") if item.get(k)]
    model = {
        "provenance": {
            "schema_version": "1.0", "dcfkit_version": E.ENGINE_VERSION, "ticker": t, "company": company,
            "producer": "import", "validation_mode": "imported",
            "created_at": legacy_created.isoformat(timespec="seconds"),
            "imported_from": {
                "path": "; ".join(files), "source": "grok", "generation": GENERATION, "fidelity": "rows",
                "source_scale": "billions" if scale != 1.0 else "millions", "convention_detected": None,
                "quality": "verified", "legacy_created": legacy_created.isoformat(timespec="seconds"),
                "notes": None,
            },
        },
        "basis": basis,
        "listing": {"symbol": t, "exchange": exchange, "quote_currency": "USD", "listed_unit": listed_unit,
                    "ordinary_per_listed_unit": ratio},
        "market": market,
        "archetype": {"kind": kind, "engine": engine_desc,
                      "decisive": {"label": decisive}},
        "drivers": list(drivers_def.values()),
        "bridge": bridge,
        "historicals": hist,
        "scenarios": scenarios,
        "headline": headline,
        "alt_valuations": alts,
        "inputs": [],
        "assumptions": assumptions,
        "guidance": [],
        "reverse_dcf": [],
        "grids": grids,
        "checks": [],
        "narrative": narrative,
        "sources": sources,
        "extra": {"legacy_summary": {
            "verdict": summ.get("verdict"), "pw_fv": pw_legacy, "pw_equity_m": summ["pw"].get("equity"),
            "pw_weights": list(weights) if weights else None, "pw_weights_source": wsrc,
            "vs_market": [(a, b) for a, b in summ["pw"].get("vs", [])],
            "summary_fv_label": fv_lab,
        }},
    }
    if listed_unit != "common":
        import_notes.append(
            f"listing: Grok values per {listed_unit} with an {listed_unit}-equivalent share count; converted to "
            f"ordinary shares (x{ratio:g} ordinary per {listed_unit}) so value_per_ordinary x {ratio:g}"
            + (f" x FX {fx:g}" if fx != 1 else "") + f" = the workbook's per-{listed_unit} value")
    if t == "SNY":
        import_notes.append("SNY: the workbook's 'FV / Ordinary (EUR)' divides by 1,210M diluted ordinary shares, "
                            "while its headline USD per-ADR value divides by 2,396.1M ADRs (= 1,198.05M "
                            "ordinary); the import reproduces the per-ADR headline (legacy EUR/ordinary kept in "
                            "results.legacy_fv_per_ordinary_eur)")
    first_end = period_for(t, years[0])[1]
    days_in = (as_of - date(first_end.year - 1, first_end.month, min(first_end.day, 28))).days
    if 30 < days_in < 365:
        import_notes.append(f"Grok discounts the full first projected year ({period_for(t, years[0])[0]}) at "
                            f"t=0.5 from the as-of date although ~{days_in} days of it had elapsed (no stub); "
                            f"kept for fidelity (stub_fraction 1)")
    if t in FISCAL and as_of >= period_for(t, years[0])[1]:
        import_notes.append(f"first projected fiscal year {period_for(t, years[0])[0]} ended "
                            f"{period_for(t, years[0])[1]} before as_of {as_of}; Grok still discounted it as a "
                            f"full future year at t=0.5 (kept as-is, stub_fraction 1)")

    # ---------------- convention detection + reproduction
    rep, conv = E.detect_convention(model, candidates=[
        {"convention": "mid", "tv_timing": "end_last"},
        {"convention": "end", "tv_timing": "end_last"},
        {"convention": "mid", "tv_timing": "mid_last"},
    ])
    conv_info = dict(conv or {})
    conv_info.update({"stub_fraction": 1.0, "grade": rep.get("grade"), "residual": rep.get("residual"),
                      "method": "dcfkit.engine.detect_convention"})
    df_check = check_discount_factors(projs, scenarios)
    if df_check:
        conv_info["sheet_discount_factors"] = df_check
    model["provenance"]["imported_from"]["convention_detected"] = conv_info
    base_ok = close(base_res["value_per_share"], num(fv_vals[1]), 0.005, 1e-9)
    if rep.get("grade") not in ("exact", "approx") or not base_ok:
        model["provenance"]["imported_from"]["quality"] = "unverified"
    model["provenance"]["imported_from"]["notes"] = "; ".join(dict.fromkeys(import_notes)) or None
    DCF.model_validate(model)
    return model


def core_label(p, key):
    rx = re.compile(CORE[key])
    for l in p["lines"]:
        if rx.search(norm(l["label"])):
            return l["label"]
    return None


def check_discount_factors(projs, scenarios):
    """Compare the sheet's discount factors with mid-year (t-0.5) at the scenario WACC."""
    out = {}
    for s in scenarios:
        p = projs[s["id"]]
        core, _ = proj_core(p)
        if "df" not in core:
            continue
        w = s["discounting"]["wacc"]
        dfs = core["df"][1]["values"]
        worst = max(abs(d - (1 + w) ** -(j + 0.5)) for j, d in enumerate(dfs))
        out[s["id"]] = {"max_abs_diff_vs_mid_year": worst}
    return out


def build_bridge(t, summ, A, bridge_by_scn, notes):
    """Top-level bridge items + per-scenario overrides when amounts differ."""
    base = bridge_by_scn["base"]
    same = all(len(v) == len(base) and all(close(a[1], b[1], 1e-12, 1e-9) for a, b in zip(v, base))
               for v in bridge_by_scn.values())
    items, overrides = [], {}

    def kind_for(label, x):
        n = norm(label)
        if "lease" in n:
            return "op_lease"
        if "equity investment" in n or "investments" in n and "equity" in n:
            return "investment_stake"
        if "net debt" in n or "net cash" in n:
            return "cash" if x > 0 else "debt"
        if "debt" in n:
            return "debt"
        if "cash" in n:
            return "cash"
        return "other"

    def make(label, x, suffix=""):
        n = norm(label)
        out = []
        if ("net debt" in n or "net cash" in n):
            split = split_net_debt(summ, A, x)
            if split:
                (cl, cash), (dl, debt) = split
                out.append({"id": "cash" + suffix, "label": f"{cl} (from net-debt line)", "kind": "cash",
                            "equity_effect": cash, "source_id": None})
                if debt:
                    out.append({"id": "debt" + suffix, "label": f"{dl} (from net-debt line)", "kind": "debt",
                                "equity_effect": -debt, "source_id": None})
                return out
        k = kind_for(label, x)
        item = {"id": slug(re.sub(r"^\([-+−]\)\s*|\(\$m\)|\(€m\)|\(\$b\)", "", label,
                                  flags=re.I))[:24].strip("_") + suffix,
                "label": label, "kind": k, "equity_effect": x}
        if k == "op_lease":
            item["debatable"] = True
            item["rationale"] = ("Grok deducts (operating) lease liabilities from equity; flip off to see "
                                 "the ex-lease value" if t != "CRWV" else
                                 "Grok's primary value deducts $10.05B operating leases ('YES (primary)'); "
                                 "the workbook also shows an ex-lease FV/share (see alt_valuations)")
        out.append(item)
        return out

    if same:
        for label, x in base:
            items += make(label, x)
    else:
        notes.append("bridge amounts differ by scenario -> per-scenario bridge items with bridge_overrides")
        for sid, its in bridge_by_scn.items():
            for label, x in its:
                for it in make(label, x, suffix=f"_{sid}"):
                    it["included"] = False
                    items.append(it)
                    overrides.setdefault(sid, {})[it["id"]] = True
    for it in items:
        if it.get("source_id") is None:
            it.pop("source_id", None)
    return items, overrides


def split_net_debt(summ, A, x):
    """Split a signed net-debt equity effect into (cash, debt) when the workbook's
    cash and debt lines foot to it exactly."""
    cand_cash, cand_debt = [], []
    rows = [(lab, v) for lab, v, _ in summ["snapshot"]]
    rows += [(str(r[0]), r[1]) for r in A.rows if r and r[0] is not None and len(r) > 1]
    for lab, v in rows:
        n = norm(lab)
        val = num(v)
        if val is None or "net" in n:
            continue
        val *= money_scale(lab)
        if re.match(r"^(total )?cash\b", n):
            cand_cash.append((lab.strip(), val))
        elif re.search(r"\bdebt\b|bank loans|term loan", n) and "weight" not in n and "cost" not in n:
            cand_debt.append((lab.strip(), val))
    for cl, c in cand_cash:
        for dl, d in cand_debt:
            if abs((c - d) - x) <= 1e-9 * max(1.0, abs(x)) + 1e-9:
                return (re.sub(r"\s*\(.*?\)\s*$", "", cl), c), (re.sub(r"\s*\(.*?\)\s*$", "", dl), d)
    return None


ALT_RULES = [
    (r"cash.?floor", "net_cash_floor"),
    (r"exit.?multiple", "exit_multiple"),
    (r"mature|mid-cycle|peak-sales annuity", "mature_state"),
    (r"eps|\bpe\b|x pe|ni check|pe check", "pe_multiple"),
]


def alt_valuations(summ, fv_lab, A, unit_ccy):
    out = []
    exit_mult = num(A.value(r"^exit ev/ebitda")[0])
    for lab, vals, _ in summ["block"]:
        n = norm(lab)
        if lab == fv_lab or n.startswith("fv / ordinary") or n.startswith("%") or n.startswith("upside"):
            continue
        if not re.search(r"fv|fair value|check|floor", n):
            continue
        vv = [num(v) for v in vals[:3]]
        if any(v is None for v in vv):
            continue
        method = next((m for rx, m in ALT_RULES if re.search(rx, n)), "other")
        note = None
        for fl, fv in summ["footnotes"]:
            if num(fv) is not None and close(num(fv), vv[1], 1e-9, 1e-9):
                note = fl.lstrip("*").strip()
        inputs = {"multiple": exit_mult, "metric": "ebitda"} if method == "exit_multiple" and exit_mult else {}
        if all(close(v, vv[0], 1e-12, 1e-12) for v in vv):
            out.append({"id": slug(lab), "label": lab, "method": method, "value_per_share": vv[0],
                        "inputs": inputs, "note": note})
        else:
            for (sid, _), v in zip(SCENARIOS, vv):
                out.append({"id": f"{slug(lab)}_{sid}", "label": f"{lab} ({sid})", "method": method,
                            "scenario_id": sid, "value_per_share": v, "inputs": inputs, "note": note})
    for a in out:
        if a.get("note") is None:
            a.pop("note")
    return out


def parse_grids(sh, base_vps, base_wacc, base_g, notes):
    if sh is None:
        return []
    grids, k = [], 0
    i = 0
    while i < len(sh.rows):
        a = sh.a(i)
        if isinstance(a, str) and "\\" in a:
            ylab, _, xlab = a.partition("\\")
            xs = []
            for v in sh.rows[i][1:]:
                x = num(v)
                if x is None:
                    break
                xs.append(x)
            ys, z = [], []
            j = i + 1
            while j < len(sh.rows) and sh.a(j) is not None:
                y = num(sh.a(j))
                if y is None:
                    break
                vals = [num(v) for v in sh.rows[j][1:1 + len(xs)]]
                ys.append(y)
                z.append(vals)
                j += 1
            title, cands = None, []
            for b in range(i - 1, max(-1, i - 6), -1):
                if isinstance(sh.a(b), str) and "\\" not in sh.a(b):
                    cands.append(sh.a(b).strip())
            for c in cands:   # prefer the section title over a subtitle line
                if re.match(r"^(sensitivity|\d\))", c, re.I) or c.isupper():
                    title = c
                    break
            title = title or (cands[0] if cands else None)
            if xs and ys:
                is_wg = re.search(r"wacc", norm(ylab)) and re.fullmatch(r"g|terminal g", norm(xlab).strip())
                gid = "wacc_g" if is_wg and not any(g["id"] == "wacc_g" for g in grids) else f"sens_{k + 1}"
                k += 1
                grid = {"id": gid, "label": (title or a)[:200], "scenario_id": "base",
                        "x": {"key": "g" if is_wg else slug(xlab), "label": xlab.strip(), "values": xs},
                        "y": {"key": "wacc" if is_wg else slug(ylab), "label": ylab.strip(), "values": ys},
                        "output": "value_ps", "z": z}
                grids.append(grid)
                if is_wg and base_wacc is not None and base_g is not None:
                    for yi, y in enumerate(ys):
                        for xi, x in enumerate(xs):
                            if close(y, base_wacc, 1e-9, 1e-9) and close(x, base_g, 1e-9, 1e-9) and \
                                    z[yi][xi] is not None and not close(z[yi][xi], base_vps, 1e-6, 1e-6):
                                notes.append(f"sensitivity base cell {z[yi][xi]} != base value {base_vps}")
            i = j
        else:
            i += 1
    return grids


def parse_historicals(t, H, unit_ccy):
    out = {"rows": [], "share_count_series": [], "extra": {}}
    if H is None:
        return out
    annual = []
    for i, r in enumerate(H.rows):
        if not (r and isinstance(r[0], str) and re.match(r"^(year|fy\b|fy \(|fiscal year)", norm(r[0]))):
            continue
        hdrs = [norm(x) if x is not None else "" for x in r]
        if not any(re.search(r"rev|sales", h) for h in hdrs[1:]):
            continue
        cols = {}
        for j, h in enumerate(hdrs[1:], 1):
            if re.search(r"^(revenue|core sales|net sales)", h):
                cols.setdefault("revenue", j)
            elif re.search(r"ebitda", h) and "margin" not in h:
                cols.setdefault("ebitda", j)
            elif h in ("fcf", "free cash flow"):
                cols.setdefault("fcf", j)
            elif h.startswith("net income"):
                cols.setdefault("net_income", j)
            elif h == "capex":
                cols.setdefault("capex", j)
        for r2 in H.rows[i + 1:]:
            if not r2 or r2[0] is None:
                break
            ym = re.search(r"(\d{4})", str(r2[0]))
            if not ym:
                break
            rec = {"label": str(r2[0]), "year": int(ym.group(1))}
            for key, j in cols.items():
                rec[key] = num(r2[j]) if j < len(r2) else None
            annual.append(rec)
    if annual:
        out["extra"]["annual"] = annual
        out["extra"]["note"] = "annual table from the Grok Historical sheet; FCF is reported (levered) FCF"
    seen = set()
    for rec in annual:
        if rec.get("fcf") is None or rec.get("revenue") is None or rec["year"] in seen:
            continue
        seen.add(rec["year"])
        mmdd, off = HIST_FISCAL.get(t, ("12-31", 0))
        y = rec["year"] + off
        pend = date(y, int(mmdd[:2]), int(mmdd[3:]))
        row = {"period": rec["label"] if rec["label"].upper().startswith("FY") else
               (f"FY{rec['year']}" if t in HIST_FISCAL else str(rec["year"])),
               "period_end": pend.isoformat(), "is_actual": True, "revenue": rec["revenue"],
               "fcff": rec["fcf"], "drivers": {}}
        if rec.get("ebitda") is not None:
            row["ebitda"] = rec["ebitda"]
        if rec.get("capex") is not None:
            row["capex"] = rec["capex"]
        out["rows"].append(row)
    out["rows"].sort(key=lambda r: r["period_end"])
    if out["rows"]:
        out["extra"]["fcff_note"] = "historical fcff = reported free cash flow (OCF - capex) from the workbook"
    if not out["extra"]:
        out.pop("extra")
    return out


def decisive_label(mtxt, book, summ):
    m = re.search(r"([A-Z][^.]{3,80}?) is the key driver", mtxt)
    if m:
        return m.group(1).strip()
    sh = book.get("Sensitivity")
    if sh is not None:
        titles = [str(a) for a in (sh.a(i) for i in range(len(sh.rows)))
                  if isinstance(a, str) and re.match(r"^(sensitivity|\d\)|[A-Z0-9])", a) and "\\" not in a
                  and ("×" in a or " x " in a.lower() or " vs " in a.lower())]
        for tt in titles:
            if not re.search(r"wacc", tt, re.I):
                tt = re.sub(r"^(SENSITIVITY\s*[—-]\s*|\d\)\s*)", "", tt).strip()
                tt = re.sub(r"\s*[—-]\s*FV.*$", "", tt)
                return tt[:120]
    kd = summ.get("key_debates")
    if kd:
        m = re.search(r"\(1\)\s*([^?(]+\??)", kd)
        if m:
            return m.group(1).strip()
    return "Base revenue path and terminal EBITDA margin"


def build_narrative(t, summ, readme, A, notes, pw, price, base_vps, decisive, ccy):
    verdict = None
    m = re.search(r"\*\*Verdict[^*]*:\*\*\s*(.+)", readme)
    if m:
        verdict = re.sub(r"\*\*", "", m.group(1)).strip()
    verdict = (verdict or summ.get("verdict") or "n/a").rstrip(" .")
    if pw is not None:
        gap = pw / price - 1
        head = (f"{verdict}: probability-weighted fair value ${pw:,.2f} vs ${price:,.2f} market ({gap:+.0%}); "
                f"base DCF ${base_vps:,.2f}.")
    else:
        head = f"{verdict}: base DCF ${base_vps:,.2f} vs ${price:,.2f} market ({base_vps / price - 1:+.0%})."
    scen = {}
    for r in A.rows:
        if r and isinstance(r[0], str) and norm(r[0]) in ("bear", "base", "bull") and isinstance(r[1], str):
            scen[norm(r[0])] = r[1].strip()
    risks = []
    for k, v in notes.items():
        if "RISK" in k.upper():
            parts = [x.strip(" .") for x in re.split(r"•|;", v) if x.strip(" .")]
            risks += parts
    change = []
    for k, v in notes.items():
        if "CHANGE" in k.upper():
            for x in re.split(r"•|;|\(\d\)", v):
                if x.strip(" ."):
                    change.append({"metric": x.strip(" .")[:200]})
    return {
        "headline": head,
        "decisive_variable": decisive,
        "thesis": summ.get("conclusion") or scen.get("base", ""),
        "bear_case": scen.get("bear", ""),
        "bull_case": scen.get("bull", ""),
        "opposing_case": summ.get("key_debates", ""),
        "change_my_mind": change[:10],
        "market_implied": "",
        "measured_vs_assumed": "Legacy Grok model: inputs hardcoded by the Grok build script; numbers not "
                               "re-verified against filings by the importer.",
        "errors_corrected": [],
        "key_risks": risks[:12],
    }


def build_sources(notes, readme):
    out = []
    txt = " ".join(v for k, v in notes.items() if "SOURCE" in k.upper())
    if not txt:
        return out
    for k, piece in enumerate([x.strip() for x in re.split(r";|\.\s", txt) if len(x.strip()) > 8][:10], 1):
        m = re.search(r"\b(10-K|10-Q|8-K|20-F|6-K)\b", piece)
        out.append({"id": f"src{k}", "kind": m.group(1) if m else "other", "description": piece[:300]})
    return out


def build_assumptions(A):
    out, seen = [], set()
    section = None
    for r in A.rows[1:]:
        if not r or r[0] is None:
            continue
        lab = str(r[0]).strip()
        v = r[1] if len(r) > 1 else None
        if v is None and lab.isupper():
            section = lab
            continue
        x = num(v) if not isinstance(v, str) or len(v) < 20 else None
        if x is None and not (isinstance(v, str) and len(v) < 40):
            continue
        aid = slug(lab)
        if aid in seen:
            continue
        seen.add(aid)
        status = "measured" if section and "CAPITAL" in section.upper() or "BALANCE" in (section or "").upper() \
            else "assumed"
        out.append({"id": aid, "label": lab, "value": x if x is not None else v, "status": status,
                    "method": f"Grok Assumptions sheet ({section})" if section else "Grok Assumptions sheet"})
    return out


def wacc_build_from(A):
    def g(pat):
        return num(A.value(pat)[0])
    wb = {"risk_free": g(r"^risk-free"), "beta": g(r"^levered beta"), "erp": g(r"^equity risk premium"),
          "cost_of_equity": g(r"^cost of equity"), "pre_tax_cost_of_debt": g(r"^pre-tax cost of debt"),
          "tax_rate": g(r"^tax rate"), "weight_equity": g(r"^equity weight"), "weight_debt": g(r"^debt weight")}
    wb = {k: v for k, v in wb.items() if v is not None}
    if not wb:
        return None
    calc = g(r"^wacc calculated")
    used = g(r"^wacc (used )?\(base\)|^wacc used")
    note = "Grok WACC build-up (Assumptions sheet)"
    if calc is not None and used is not None and not close(calc, used, 1e-9, 1e-9):
        note += f"; calculated {calc:.4f} but Base uses {used:.4f} (judgement overlay)"
    wb["note"] = note
    return wb


# --------------------------------------------------------------------------
# import into the store
# --------------------------------------------------------------------------

def _sha256(p):
    h = hashlib.sha256()
    with open(p, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def legacy_files(item):
    files = [item["xlsx"]] + [item[k] for k in ("readme", "csv", "build_py") if item.get(k)]
    if item["ticker"] in PSYCHEDELIC and (GROK_DIR / "psychedelic_dcf_engine.py").exists():
        files.append(GROK_DIR / "psychedelic_dcf_engine.py")
    return [Path(f) for f in files]


def existing_version(ticker, xlsx_rel):
    tdir = MODELS / ticker.upper()
    if not tdir.is_dir():
        return None
    for rj in sorted(tdir.glob("*/run.json")):
        meta = read_json(rj, {}) or {}
        leg = meta.get("legacy") or {}
        files = leg.get("files") or []
        if leg.get("legacy_id") == f"grok:{ticker.upper()}" or xlsx_rel in files:
            return rj.parent
    return None


def _vid_holders(vid):
    return sorted(p.parent.name for p in MODELS.glob(f"*/{vid}") if p.is_dir())


def _free_ts(ticker, ts):
    """Version ids must be globally unique (store.find_version_dir / the API look a
    vid up across all tickers), but several Grok workbooks share an mtime second.
    Bump by whole seconds until no OTHER ticker holds the id. legacy mtime itself is
    kept in run.json legacy.mtime and provenance.imported_from.legacy_created."""
    from datetime import timedelta
    while [h for h in _vid_holders(store.make_version_id("grok", "import", ts)) if h != ticker.upper()]:
        ts = ts + timedelta(seconds=1)
    return ts


def _dedupe_prior(ticker, path):
    """Rename an earlier import whose vid collides with another ticker's (keeps the
    alphabetically-first holder on the original id). Returns the (possibly new) path."""
    holders = _vid_holders(path.name)
    if len(holders) < 2 or holders[0] == ticker.upper():
        return path
    stamp = datetime.strptime(path.name[:15], "%Y%m%d-%H%M%S")
    from datetime import timedelta
    ts = stamp + timedelta(seconds=1)
    ts = _free_ts(ticker, ts)
    vid = store.make_version_id("grok", "import", ts)
    new = store.version_dir(ticker, vid)
    if new.exists():
        return path
    path.rename(new)          # rename (not delete) our own earlier output
    update_run(new, id=vid)
    d = read_json(new / "dcf.json")
    if isinstance(d, dict):
        d["provenance"]["version_id"] = vid
        write_json(new / "dcf.json", d)
    return new


def report_md(item, model, crosscheck):
    t = item["ticker"]
    imp = model["provenance"]["imported_from"]
    conv = imp.get("convention_detected") or {}
    h = model["headline"]
    lines = [
        f"<!-- imported by server.importer.grok_xlsx; legacy README reproduced verbatim below -->",
        f"# {t} -- imported Grok DCF (legacy, {model['basis']['as_of']})",
        "",
        f"- **Source:** {imp['path']} (legacy mtime {imp['legacy_created']})",
        f"- **Generation:** {imp['generation']} -- fidelity `{imp['fidelity']}`, flows in reporting-currency "
        f"millions ({model['basis']['reporting_currency']})",
        f"- **Convention:** {conv.get('convention')} / TV {conv.get('tv_timing')}, stub 1 -- engine reproduction "
        f"grade **{conv.get('grade')}** (residual {conv.get('residual')})",
        f"- **Headline ({h['method']}):** {h['value_per_share']:.4f} per {model['listing']['listed_unit']} "
        f"(base {crosscheck['imported_vps']:.4f}; legacy base {crosscheck['legacy_vps']:.4f}; "
        f"match {crosscheck['match']})",
    ]
    if imp.get("notes"):
        lines.append(f"- **Import notes:** {imp['notes']}")
    lines += ["", "---", ""]
    return "\n".join(lines) + read_readme(item)


def crosscheck_of(item, model):
    summ = parse_summary(load_book(item["xlsx"])["Summary"])
    _, fv = block_row(summ, r"^fair value /")
    legacy_vps = num(fv[1])
    base = E.scenario_by_role(model, "base")
    imported = base["results"]["value_per_share"]
    rep = E.reproduce(model)
    eng = E.value_model(copy.deepcopy(model))
    eng_base = E.scenario_by_role(eng, "base")["results"]["value_per_share_raw"]
    pw_legacy = summ["pw"].get("fv")
    return {
        "legacy_vps": legacy_vps, "imported_vps": imported,
        "engine_vps": eng_base,
        "match": bool(close(legacy_vps, imported, 0.005, 1e-9) and close(legacy_vps, eng_base, 0.005, 1e-9)),
        "pw_legacy": pw_legacy, "pw_imported": model["headline"].get("value_per_share"),
        "pw_match": bool(pw_legacy is None or close(pw_legacy, model["headline"].get("value_per_share"),
                                                    1e-6, 1e-9)),
        "grade": rep["grade"], "residual": rep.get("residual"),
    }


def import_one(item, *, force=False):
    t = item["ticker"]
    xlsx_rel = rel(item["xlsx"])
    prior = existing_version(t, xlsx_rel)
    if prior is not None and not force:
        meta = read_json(prior / "run.json", {}) or {}
        return {"ticker": t, "status": "skipped", "version_id": prior.name, "reason": "already imported",
                "run_status": meta.get("status"), "crosscheck": (meta.get("legacy") or {}).get("crosscheck")}
    model = convert(item)
    cc = crosscheck_of(item, model)
    files = legacy_files(item)
    legacy = {
        "legacy_id": f"grok:{t}", "files": [rel(f) for f in files], "source": "grok",
        "generation": GENERATION, "main": xlsx_rel,
        "mtime": item["mtime"].isoformat(timespec="seconds"),
        "sha256": {rel(f): _sha256(f) for f in files},
        "crosscheck": {"legacy_vps": cc["legacy_vps"], "imported_vps": cc["imported_vps"], "match": cc["match"],
                       "engine_vps": cc["engine_vps"], "pw_legacy": cc["pw_legacy"],
                       "pw_imported": cc["pw_imported"], "pw_match": cc["pw_match"],
                       "grade": cc["grade"], "residual": cc["residual"]},
    }
    started = now_iso()
    if prior is not None:
        prior = _dedupe_prior(t, prior)
        path, vid = prior, prior.name
        update_run(path, status="running", started_at=started, finished_at=None, error=None, legacy=legacy,
                   kind="import")
    else:
        vid, path = store.create_version(t, "grok", "import", ts=_free_ts(t, item["mtime"]), vendor=False,
                                         run_meta={"kind": "import", "legacy": legacy, "status": "running",
                                                   "started_at": started,
                                                   "instructions": "deterministic import: "
                                                                   "server.importer.grok_xlsx"})
    try:
        ldir = path / "legacy"
        ldir.mkdir(exist_ok=True)
        for f in files:
            shutil.copy2(f, ldir / f.name)     # copy, never move
        model["provenance"]["version_id"] = vid
        write_json(path / "dcf.json", model)
        (path / "report.md").write_text(report_md(item, model, cc))
        update_run(path, status="validating")
        from server.validation import run_validation
        val = run_validation(path, "imported")
        status = "valid" if val.get("ok") else "invalid"
        update_run(path, status=status, finished_at=now_iso(),
                   error=None if val.get("ok") else "validation failures (imported mode)")
    except Exception as exc:
        update_run(path, status="failed", finished_at=now_iso(), error=f"{type(exc).__name__}: {exc}")
        raise
    return {"ticker": t, "status": status, "version_id": vid, "validation": val.get("summary"),
            "reproduction": (val.get("reproduction") or {}).get("grade"), "crosscheck": legacy["crosscheck"],
            "notes": model["provenance"]["imported_from"].get("notes"),
            "convention": model["provenance"]["imported_from"].get("convention_detected")}


def _report_entry(res, model_notes=None):
    cc = res.get("crosscheck") or {}
    conv = res.get("convention") or {}
    return {
        "status": res.get("status"), "version_id": res.get("version_id"),
        "base_vps_legacy": cc.get("legacy_vps"), "base_vps_imported": cc.get("imported_vps"),
        "base_vps_engine": cc.get("engine_vps"), "base_match": cc.get("match"),
        "pw": {"legacy": cc.get("pw_legacy"), "imported": cc.get("pw_imported"), "match": cc.get("pw_match")},
        "grade": cc.get("grade"), "residual": cc.get("residual"),
        "convention": {k: conv.get(k) for k in ("convention", "tv_timing", "stub_fraction")} if conv else None,
        "validation": res.get("validation"), "notes": res.get("notes") or res.get("reason") or res.get("error"),
    }


def import_tickers(tickers=None, force=False):
    """Import the given tickers (None/empty = all). Updates data/imports/grok_report.json.

    Returns {"results": {ticker: {...}}, "summary": {...}}.
    """
    inv = {it["ticker"]: it for it in inventory()}
    want = [x.upper() for x in tickers] if tickers else sorted(inv)
    results = {}
    for t in want:
        it = inv.get(t)
        if it is None:
            results[t] = {"ticker": t, "status": "failed", "error": "not in the Grok inventory"
                          + (" (CRBS is a duplicate of CBRS)" if t in SKIP else "")}
            continue
        try:
            results[t] = import_one(it, force=force)
        except Exception as exc:
            results[t] = {"ticker": t, "status": "failed", "error": f"{type(exc).__name__}: {exc}"}
    report = read_json(REPORT_PATH, {}) or {}
    per = report.get("tickers") or {}
    for t, res in results.items():
        if res.get("status") == "skipped" and t in per:
            continue
        per[t] = _report_entry(res)
    summ = {"total": len(per), "imported": 0, "exact": 0, "approx": 0, "none": 0, "failed": 0,
            "valid": 0, "invalid": 0, "base_match": 0, "pw_match": 0}
    for e in per.values():
        if e["status"] in ("valid", "invalid", "skipped"):
            summ["imported"] += 1
        if e["status"] == "failed":
            summ["failed"] += 1
        if e["status"] in ("valid", "invalid"):
            summ[e["status"]] += 1
        if e.get("grade") in ("exact", "approx", "none"):
            summ[e["grade"]] += 1
        summ["base_match"] += bool(e.get("base_match"))
        summ["pw_match"] += bool((e.get("pw") or {}).get("match"))
    report = {"generated_at": now_iso(), "importer": "server.importer.grok_xlsx", "legacy_dir": rel(GROK_DIR),
              "skipped_duplicates": sorted(SKIP), "summary": summ, "tickers": dict(sorted(per.items()))}
    write_json(REPORT_PATH, report)
    store.STORE.invalidate()
    return {"results": results, "summary": summ}


def main(argv=None):
    ap = argparse.ArgumentParser(description="Import legacy Grok DCF workbooks into DCF Studio")
    ap.add_argument("--all", action="store_true", help="import every Grok workbook")
    ap.add_argument("--tickers", nargs="+", help="only these tickers")
    ap.add_argument("--force", action="store_true", help="re-convert even if already imported")
    ap.add_argument("--dry-run", action="store_true", help="convert + validate schema only; write nothing")
    a = ap.parse_args(argv)
    if not a.all and not a.tickers:
        ap.error("pass --all or --tickers")
    if a.dry_run:
        inv = {it["ticker"]: it for it in inventory()}
        for t in (a.tickers or sorted(inv)):
            try:
                m = convert(inv[t.upper()])
                c = m["provenance"]["imported_from"]["convention_detected"]
                print(f"{t:5} OK grade={c['grade']} residual={c['residual']}")
            except Exception as exc:
                print(f"{t:5} FAIL {type(exc).__name__}: {exc}")
        return 0
    out = import_tickers(None if a.all and not a.tickers else a.tickers, force=a.force)
    for t, r in out["results"].items():
        cc = r.get("crosscheck") or {}
        print(f"{t:5} {r['status']:8} {r.get('version_id') or '':32} grade={cc.get('grade')} "
              f"base {cc.get('legacy_vps')} -> {cc.get('imported_vps')} {r.get('error') or ''}")
    print(json.dumps(out["summary"]))
    return 0 if not any(r["status"] == "failed" for r in out["results"].values()) else 1


if __name__ == "__main__":
    sys.exit(main())
